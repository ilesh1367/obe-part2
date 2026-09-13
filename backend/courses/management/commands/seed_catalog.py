from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand
from django.db import transaction
from openpyxl import load_workbook

from courses.models import AcademicSession, FacultyProfile, NbaSubjectCatalog

NBA_FILE = 'NBA codes-ODD 2026.xlsx'
DATA_FILE = 'data.xlsx'


class Command(BaseCommand):
    help = 'Seed AcademicSession, NBA catalog (2026 Odd), and faculty directory from data/*.xlsx'

    def add_arguments(self, parser):
        parser.add_argument('--year', type=int, default=2026)
        parser.add_argument('--semester', default='ODD', choices=['ODD', 'EVEN'])

    def handle(self, *args, **options):
        data_dir = Path(settings.BASE_DIR).parent / 'data'
        if not data_dir.exists():
            data_dir = Path(settings.BASE_DIR) / 'data'
        year = options['year']
        sem = options['semester']
        session, _ = AcademicSession.objects.get_or_create(calendar_year=year, semester_type=sem)
        self.stdout.write(f'Session: {session.label}')

        nba_path = data_dir / NBA_FILE
        if nba_path.exists():
            created, updated = self._seed_nba(session, nba_path)
            self.stdout.write(self.style.SUCCESS(f'NBA catalog: {created} created, {updated} updated'))
        else:
            self.stdout.write(self.style.WARNING(f'Missing {nba_path}'))

        faculty_path = data_dir / DATA_FILE
        if faculty_path.exists():
            created = self._seed_faculty(faculty_path)
            self.stdout.write(self.style.SUCCESS(f'Faculty directory: {created} created'))
        else:
            self.stdout.write(self.style.WARNING(f'Missing {faculty_path}'))

    @transaction.atomic
    def _seed_nba(self, session, path):
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        created = updated = 0
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                continue
            if not row or not row[2] or not row[6]:
                continue
            program = str(row[1] or '').strip()
            course_code = str(row[2]).strip()
            course_name = str(row[3] or '').strip()
            year_of_study = int(row[4] or 0)
            semester_number = int(row[5] or 0)
            nba_code = str(row[6]).strip()
            if not course_code or not nba_code:
                continue
            obj, was_created = NbaSubjectCatalog.objects.update_or_create(
                session=session,
                nba_code=nba_code,
                course_code=course_code,
                defaults={
                    'program_name': program,
                    'course_name': course_name,
                    'year_of_study': year_of_study,
                    'semester_number': semester_number,
                },
            )
            if was_created:
                created += 1
            else:
                updated += 1
        return created, updated

    @transaction.atomic
    def _seed_faculty(self, path):
        wb = load_workbook(path, data_only=True)
        if 'Faculty List' not in wb.sheetnames:
            return 0
        ws = wb['Faculty List']
        created = 0
        for row in ws.iter_rows(values_only=True):
            name = str(row[0] or '').strip()
            if not name:
                continue
            _, was_created = FacultyProfile.objects.get_or_create(
                full_name=name,
                campus=FacultyProfile.Campus.UNKNOWN,
                defaults={'is_active': True},
            )
            if was_created:
                created += 1
        return created
